from pathlib import Path
import sys

import fitz
from docx import Document


def main() -> None:
    output = Path(sys.argv[1])
    marker = sys.argv[2]
    output.mkdir(parents=True, exist_ok=True)

    (output / "enterprise-guide.txt").write_text(
        (f"企业专属编号 {marker}。系统支持用友 U8 Cloud 标准 API 对接。" * 30),
        encoding="utf-8",
    )

    document = Document()
    document.add_heading("交付方案", level=1)
    document.add_paragraph("企业项目采用分阶段实施与验收方案。" * 30)
    document.save(output / "delivery-plan.docx")

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Enterprise service and API integration guide. " * 20)
    pdf.save(output / "service-guide.pdf")
    pdf.close()


if __name__ == "__main__":
    main()
