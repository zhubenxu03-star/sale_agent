from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from openpyxl import load_workbook

from app.core.config import settings
from app.services.knowledge.errors import KnowledgeProcessingError


@dataclass(slots=True)
class DocumentSection:
    text: str
    page_number: int | None = None
    sheet_name: str | None = None
    row_start: int | None = None
    row_end: int | None = None
    section_title: str | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


def _decode_text(data: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise KnowledgeProcessingError("TEXT_ENCODING_UNSUPPORTED", "无法识别文本文件编码")


def parse_pdf(path: Path) -> list[DocumentSection]:
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise KnowledgeProcessingError("PDF_PARSE_FAILED", "PDF文件无法解析") from exc
    sections: list[DocumentSection] = []
    title = str(document.metadata.get("title") or "").strip() or None
    for index, page in enumerate(document, start=1):
        text = page.get_text("text").strip()
        if text:
            sections.append(
                DocumentSection(
                    text=text,
                    page_number=index,
                    section_title=title,
                    metadata={"pdf_title": title} if title else {},
                )
            )
    document.close()
    if sum(len(item.text) for item in sections) < 30:
        raise KnowledgeProcessingError(
            "OCR_REQUIRED", "该PDF可能是扫描文件，当前版本暂不支持OCR识别。"
        )
    return sections


def parse_docx(path: Path) -> list[DocumentSection]:
    try:
        document = Document(path)
    except Exception as exc:
        raise KnowledgeProcessingError("DOCX_PARSE_FAILED", "DOCX文件无法解析") from exc
    sections: list[DocumentSection] = []
    current_title: str | None = None
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.lower().startswith("heading") or style.startswith("标题"):
            current_title = text
            continue
        sections.append(
            DocumentSection(text=text, section_title=current_title, metadata={"style": style})
        )
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            sections.append(
                DocumentSection(
                    text="\n".join(rows),
                    section_title=current_title or f"表格 {table_index}",
                    metadata={"table_index": table_index},
                )
            )
    return sections


def parse_text(path: Path, markdown: bool = False) -> list[DocumentSection]:
    text, encoding = _decode_text(path.read_bytes())
    sections: list[DocumentSection] = []
    current_title: str | None = None
    blocks: list[str] = []

    def flush() -> None:
        if blocks:
            value = "\n".join(blocks).strip()
            if value:
                sections.append(
                    DocumentSection(
                        text=value,
                        section_title=current_title,
                        metadata={"encoding": encoding},
                    )
                )
            blocks.clear()

    for line in text.splitlines():
        stripped = line.strip()
        heading = re.match(r"^#{1,6}\s+(.+)$", stripped) if markdown else None
        if heading:
            flush()
            current_title = heading.group(1).strip()
        elif not stripped:
            flush()
        else:
            blocks.append(stripped)
    flush()
    return sections


def parse_xlsx(path: Path) -> list[DocumentSection]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise KnowledgeProcessingError("XLSX_PARSE_FAILED", "XLSX文件无法解析") from exc
    sections: list[DocumentSection] = []
    total_cells = 0
    try:
        for sheet in workbook.worksheets:
            if sheet.max_row > settings.knowledge_xlsx_max_rows:
                raise KnowledgeProcessingError("XLSX_TOO_MANY_ROWS", "工作表行数超过限制")
            if sheet.max_column > settings.knowledge_xlsx_max_columns:
                raise KnowledgeProcessingError("XLSX_TOO_MANY_COLUMNS", "工作表列数超过限制")
            rows = sheet.iter_rows(values_only=True)
            header: list[str] | None = None
            for row_number, row in enumerate(rows, start=1):
                values = ["" if value is None else str(value).strip() for value in row]
                total_cells += len(values)
                if total_cells > settings.knowledge_xlsx_max_cells:
                    raise KnowledgeProcessingError(
                        "XLSX_TOO_MANY_CELLS", "工作簿单元格总量超过限制"
                    )
                if not any(values):
                    continue
                if header is None:
                    header = [value or f"列{index + 1}" for index, value in enumerate(values)]
                    continue
                pairs = [
                    f"{header[index] if index < len(header) else f'列{index + 1}'}: {value}"
                    for index, value in enumerate(values)
                    if value
                ]
                if pairs:
                    sections.append(
                        DocumentSection(
                            text="；".join(pairs),
                            sheet_name=sheet.title,
                            row_start=row_number,
                            row_end=row_number,
                            section_title=sheet.title,
                        )
                    )
    finally:
        workbook.close()
    return sections


def parse_csv(path: Path) -> list[DocumentSection]:
    text, encoding = _decode_text(path.read_bytes())
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(text.splitlines(), dialect)
    sections: list[DocumentSection] = []
    header: list[str] | None = None
    for row_number, row in enumerate(reader, start=1):
        if row_number > settings.knowledge_csv_max_rows:
            raise KnowledgeProcessingError("CSV_TOO_MANY_ROWS", "CSV行数超过限制")
        values = [value.strip() for value in row]
        if not any(values):
            continue
        if header is None:
            header = [value or f"列{index + 1}" for index, value in enumerate(values)]
            continue
        pairs = [
            f"{header[index] if index < len(header) else f'列{index + 1}'}: {value}"
            for index, value in enumerate(values)
            if value
        ]
        if pairs:
            sections.append(
                DocumentSection(
                    text="；".join(pairs),
                    row_start=row_number,
                    row_end=row_number,
                    section_title="CSV数据",
                    metadata={"encoding": encoding, "delimiter": dialect.delimiter},
                )
            )
    return sections


def parse_document(path: Path, extension: str) -> list[DocumentSection]:
    parsers = {
        "pdf": parse_pdf,
        "docx": parse_docx,
        "txt": lambda value: parse_text(value, False),
        "md": lambda value: parse_text(value, True),
        "xlsx": parse_xlsx,
        "csv": parse_csv,
    }
    try:
        sections = parsers[extension](path)
    except KeyError as exc:
        raise KnowledgeProcessingError("UNSUPPORTED_FILE_TYPE", "不支持的文件类型") from exc
    if not sections or not any(item.text.strip() for item in sections):
        raise KnowledgeProcessingError("EMPTY_DOCUMENT", "文档中没有可处理的有效文本")
    return sections
