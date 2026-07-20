from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any


@dataclass
class ParsedRecord:
    values: dict[str, Any]
    index: int


@dataclass
class ParsedDocument:
    records: list[ParsedRecord]
    headers: list[str]
    format: str


LINE_RE = re.compile(r"^\s*\[?(?P<timestamp>[^\]]{8,32})\]?\s*(?P<sender>[^:：]{1,40})[:：]\s*(?P<content>.+)$")


def _timestamp(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value).strip()


def parse_file(path: Path, extension: str, max_records: int | None = None) -> ParsedDocument:
    ext = extension.lower().lstrip(".")
    if ext in {"txt", "md"}:
        records: list[ParsedRecord] = []
        for idx, line in enumerate(path.read_text(encoding="utf-8-sig", errors="replace").splitlines()):
            match = LINE_RE.match(line)
            if match:
                values = {
                    "conversation_id": None,
                    "sender_name": match.group("sender").strip(),
                    "sender_role": match.group("sender").strip(),
                    "content": match.group("content").strip(),
                    "timestamp": match.group("timestamp").strip(),
                }
            elif line.strip():
                values = {"content": line.strip(), "sender_role": None, "sender_name": None}
            else:
                continue
            records.append(ParsedRecord(values, idx))
        return ParsedDocument(records[:max_records] if max_records else records, [], ext)
    if ext == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            rows = [ParsedRecord(dict(row), idx) for idx, row in enumerate(reader)]
        return ParsedDocument(rows[:max_records] if max_records else rows, list(reader.fieldnames or []), ext)
    if ext == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in (rows[0] if rows else [])]
        records = [
            ParsedRecord({header: value for header, value in zip(headers, row, strict=False) if header}, idx)
            for idx, row in enumerate(rows[1:])
            if any(value not in (None, "") for value in row)
        ]
        return ParsedDocument(records[:max_records] if max_records else records, headers, ext)
    if ext == "json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        records: list[ParsedRecord] = []
        for conversation in payload.get("conversations", payload if isinstance(payload, list) else []):
            if not isinstance(conversation, dict):
                continue
            base = {key: conversation.get(key) for key in ("conversation_id", "industry", "outcome", "deal_amount", "sales_stage")}
            for message in conversation.get("messages", []):
                if isinstance(message, dict):
                    values = {**base, **message}
                    records.append(ParsedRecord(values, len(records)))
        return ParsedDocument(records[:max_records] if max_records else records, [], ext)
    if ext == "docx":
        from docx import Document

        records = []
        for idx, paragraph in enumerate(Document(path).paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            match = LINE_RE.match(text)
            records.append(
                ParsedRecord(
                    {
                        "sender_name": match.group("sender").strip() if match else None,
                        "sender_role": match.group("sender").strip() if match else None,
                        "content": match.group("content").strip() if match else text,
                        "timestamp": match.group("timestamp").strip() if match else None,
                    },
                    idx,
                )
            )
        return ParsedDocument(records[:max_records] if max_records else records, [], ext)
    raise ValueError(f"unsupported champion file format: {ext}")
