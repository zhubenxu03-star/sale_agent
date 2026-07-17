import csv
import math
from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook
from pydantic import ValidationError

from app.core.config import Settings
from app.core.exceptions import AppException
from app.schemas.knowledge import KnowledgeBaseUpdate, KnowledgeSearchRequest
from app.services.knowledge.chunker import build_chunks, clean_text, estimate_tokens
from app.services.knowledge.embeddings import (
    DeterministicTestEmbeddingProvider,
    EmbeddingProvider,
    get_embedding_provider,
)
from app.services.knowledge.errors import KnowledgeProcessingError
from app.services.knowledge.parser import DocumentSection, parse_document
from app.services.knowledge.storage import sanitize_filename, validate_file_format


def long_text(marker: str = "企业知识") -> str:
    return (f"{marker}用于产品、价格、交付和常见问题管理。" * 20) + "支持用友U8 Cloud接口对接。"


def test_deterministic_embedding_is_repeatable() -> None:
    provider = DeterministicTestEmbeddingProvider(1536)
    assert (
        provider.embed_texts(["用友U8 Cloud对接"])[0]
        == provider.embed_texts(["用友U8 Cloud对接"])[0]
    )


def test_deterministic_embedding_is_normalized() -> None:
    vector = DeterministicTestEmbeddingProvider(1536).embed_texts(["企业知识库"])[0]
    assert math.isclose(sum(value * value for value in vector), 1.0, rel_tol=1e-6)


def test_embedding_count_mismatch_fails() -> None:
    provider = DeterministicTestEmbeddingProvider(1536)
    with pytest.raises(KnowledgeProcessingError, match="向量数量"):
        provider.validate(["a"], [])


def test_embedding_dimension_mismatch_fails() -> None:
    provider = DeterministicTestEmbeddingProvider(1536)
    with pytest.raises(KnowledgeProcessingError, match="向量维度"):
        provider.validate(["a"], [[0.0] * 3])


def test_production_rejects_test_embedding() -> None:
    with pytest.raises(ValidationError, match="production cannot use"):
        Settings(_env_file=None, app_env="production", embedding_provider="test")


def test_production_requires_embedding_credentials() -> None:
    with pytest.raises(ValidationError, match="configuration is incomplete"):
        Settings(
            _env_file=None,
            app_env="production",
            embedding_provider="openai_compatible",
        )


def test_development_returns_test_provider() -> None:
    config = Settings(_env_file=None, app_env="development", embedding_provider="test")
    assert get_embedding_provider(config).mode == "test"


def test_blank_search_query_is_rejected() -> None:
    with pytest.raises(ValidationError):
        KnowledgeSearchRequest(query="   ")


def test_null_knowledge_base_name_is_rejected() -> None:
    with pytest.raises(ValidationError):
        KnowledgeBaseUpdate(name=None)


def test_chunker_keeps_source_metadata() -> None:
    chunks = build_chunks(
        [DocumentSection(text=long_text(), page_number=8, section_title="对接能力")]
    )
    assert chunks[0].page_number == 8
    assert chunks[0].section_title == "对接能力"


def test_chunker_deduplicates_identical_content() -> None:
    text = long_text()
    chunks = build_chunks([DocumentSection(text=text), DocumentSection(text=text)])
    assert len({item.content_hash for item in chunks}) == len(chunks)


def test_chunker_splits_long_paragraph_with_overlap() -> None:
    chunks = build_chunks(
        [DocumentSection(text="".join(f"第{index}个完整句子。" for index in range(300)))],
        target_size=200,
        overlap=30,
        min_size=20,
    )
    assert len(chunks) > 2
    assert all(len(item.content) <= 230 for item in chunks)


def test_clean_text_removes_null_and_excess_whitespace() -> None:
    assert clean_text("甲\x00   乙\n\n\n丙") == "甲 乙\n\n丙"


def test_token_estimator_returns_positive_count() -> None:
    assert estimate_tokens("企业知识库") > 0


def test_filename_is_sanitized_against_traversal() -> None:
    assert sanitize_filename("../../危险<文件>.txt") == "危险_文件_.txt"


def test_fake_pdf_signature_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "fake.pdf"
    path.write_bytes(b"not-a-pdf")
    with pytest.raises(AppException) as exc:
        validate_file_format(path, "pdf")
    assert exc.value.status_code == 415


def test_invalid_office_archive_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "fake.docx"
    path.write_bytes(b"not-a-zip")
    with pytest.raises(AppException) as exc:
        validate_file_format(path, "docx")
    assert exc.value.error_code == "INVALID_OFFICE_FORMAT"


def test_binary_text_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "bad.txt"
    path.write_bytes(b"hello\x00world")
    with pytest.raises(AppException) as exc:
        validate_file_format(path, "txt")
    assert exc.value.error_code == "INVALID_TEXT_FORMAT"


def test_parse_txt_and_utf8_bom(tmp_path: Path) -> None:
    path = tmp_path / "guide.txt"
    path.write_bytes(("\ufeff" + long_text()).encode())
    sections = parse_document(path, "txt")
    assert "用友U8 Cloud" in sections[0].text
    assert sections[0].metadata["encoding"] == "utf-8-sig"


def test_parse_markdown_preserves_heading(tmp_path: Path) -> None:
    path = tmp_path / "guide.md"
    path.write_text(f"# 对接能力\n\n{long_text()}", encoding="utf-8")
    sections = parse_document(path, "md")
    assert sections[0].section_title == "对接能力"


def test_parse_docx_preserves_heading_and_table(tmp_path: Path) -> None:
    path = tmp_path / "guide.docx"
    document = Document()
    document.add_heading("交付说明", level=1)
    document.add_paragraph(long_text())
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "产品"
    table.cell(0, 1).text = "能力"
    table.cell(1, 0).text = "ERP"
    table.cell(1, 1).text = "标准接口"
    document.save(path)
    sections = parse_document(path, "docx")
    assert sections[0].section_title == "交付说明"
    assert any("ERP" in item.text for item in sections)


def test_parse_xlsx_preserves_sheet_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "price.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "价格表"
    sheet.append(["产品", "价格"])
    sheet.append(["标准版", 10000])
    workbook.save(path)
    sections = parse_document(path, "xlsx")
    assert sections[0].sheet_name == "价格表"
    assert sections[0].row_start == 2
    assert "标准版" in sections[0].text


def test_parse_csv_detects_delimiter_and_rows(tmp_path: Path) -> None:
    path = tmp_path / "cases.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as output:
        writer = csv.writer(output, delimiter=";")
        writer.writerow(["客户", "案例"])
        writer.writerow(["云启", long_text("零售案例")])
    sections = parse_document(path, "csv")
    assert sections[0].row_start == 2
    assert sections[0].metadata["delimiter"] == ";"


def test_parse_text_pdf_preserves_page(tmp_path: Path) -> None:
    path = tmp_path / "guide.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "ERP integration U8 Cloud standard API " * 8)
    document.save(path)
    document.close()
    sections = parse_document(path, "pdf")
    assert sections[0].page_number == 1


def test_scanned_pdf_requires_ocr(tmp_path: Path) -> None:
    path = tmp_path / "scan.pdf"
    document = fitz.open()
    document.new_page()
    document.save(path)
    document.close()
    with pytest.raises(KnowledgeProcessingError) as exc:
        parse_document(path, "pdf")
    assert exc.value.code == "OCR_REQUIRED"


class WrongDimensionProvider(EmbeddingProvider):
    name = "wrong"
    mode = "test"

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[0.0] * 3 for _ in texts]


def test_wrong_dimension_provider_contract() -> None:
    provider = WrongDimensionProvider(1536)
    with pytest.raises(KnowledgeProcessingError):
        provider.validate(["text"], provider.embed_texts(["text"]))
