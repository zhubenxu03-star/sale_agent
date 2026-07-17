from __future__ import annotations

import os
from collections.abc import Generator
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID, uuid4

import fitz
import psycopg
import pytest
from alembic.config import Config
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, func, select, text
from sqlalchemy.orm import Session, sessionmaker

from alembic import command
from app.core.config import settings
from app.core.security import create_access_token, hash_password
from app.db.session import get_db
from app.main import app
from app.models.knowledge import (
    DocumentStatus,
    KnowledgeBase,
    KnowledgeChunk,
    KnowledgeDocument,
    KnowledgeProcessingJob,
    KnowledgeRetrievalLog,
)
from app.models.tenant import Tenant
from app.models.user import User, UserRole
from app.services.knowledge.embeddings import (
    DeterministicTestEmbeddingProvider,
    EmbeddingProvider,
)
from app.services.knowledge.processing import process_knowledge_document

POSTGRES_URL = os.getenv(
    "TEST_DATABASE_URL",
    "postgresql+psycopg://postgres:postgres@localhost:5432/sales_agent_test",
)


@pytest.fixture(scope="session")
def postgres_factory() -> Generator[sessionmaker[Session], None, None]:
    database_name = POSTGRES_URL.rsplit("/", 1)[-1]
    maintenance_url = POSTGRES_URL.rsplit("/", 1)[0] + "/postgres"
    with psycopg.connect(
        maintenance_url.replace("postgresql+psycopg", "postgresql"), autocommit=True
    ) as connection:
        exists = connection.execute(
            "SELECT 1 FROM pg_database WHERE datname = %s", (database_name,)
        ).fetchone()
        if exists is None:
            connection.execute(f'CREATE DATABASE "{database_name}"')

    alembic_config = Config(str(Path(__file__).parents[1] / "alembic.ini"))
    alembic_config.set_main_option("script_location", str(Path(__file__).parents[1] / "alembic"))
    alembic_config.set_main_option("sqlalchemy.url", POSTGRES_URL)
    original_database_url = settings.database_url
    settings.database_url = POSTGRES_URL
    try:
        command.upgrade(alembic_config, "head")
    finally:
        settings.database_url = original_database_url

    engine = create_engine(POSTGRES_URL, pool_pre_ping=True)
    factory = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)
    yield factory
    engine.dispose()


@pytest.fixture
def postgres_client(
    postgres_factory: sessionmaker[Session], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> Generator[tuple[TestClient, sessionmaker[Session]], None, None]:
    with postgres_factory() as db:
        db.execute(text("TRUNCATE TABLE tenants CASCADE"))
        db.commit()

    def override_get_db() -> Generator[Session, None, None]:
        db = postgres_factory()
        try:
            yield db
        finally:
            db.close()

    previous_override = app.dependency_overrides.get(get_db)
    app.dependency_overrides[get_db] = override_get_db
    monkeypatch.setattr(settings, "knowledge_storage_path", tmp_path / "knowledge")
    monkeypatch.setattr(
        "app.api.v1.knowledge.process_document_task.delay",
        lambda *_: SimpleNamespace(id=f"test-{uuid4()}"),
    )
    with TestClient(app, raise_server_exceptions=False) as client:
        yield client, postgres_factory
    if previous_override is None:
        app.dependency_overrides.pop(get_db, None)
    else:
        app.dependency_overrides[get_db] = previous_override


def seed_enterprises(factory: sessionmaker[Session]) -> dict[str, object]:
    with factory() as db:
        tenant_a = Tenant(name="甲企业", code=f"tenant-a-{uuid4().hex[:8]}")
        tenant_b = Tenant(name="乙企业", code=f"tenant-b-{uuid4().hex[:8]}")
        db.add_all([tenant_a, tenant_b])
        db.flush()
        password = hash_password("Test123456")
        admin = User(
            tenant_id=tenant_a.id,
            name="管理员",
            email="admin@example.com",
            password_hash=password,
            role=UserRole.ADMIN,
        )
        manager = User(
            tenant_id=tenant_a.id,
            name="经理",
            email="manager@example.com",
            password_hash=password,
            role=UserRole.MANAGER,
        )
        sales = User(
            tenant_id=tenant_a.id,
            name="销售",
            email="sales@example.com",
            password_hash=password,
            role=UserRole.SALES,
        )
        other_admin = User(
            tenant_id=tenant_b.id,
            name="乙管理员",
            email="admin@example.com",
            password_hash=password,
            role=UserRole.ADMIN,
        )
        db.add_all([admin, manager, sales, other_admin])
        db.flush()
        base_a = KnowledgeBase(
            tenant_id=tenant_a.id,
            name="企业知识库",
            description="甲企业资料",
            created_by_user_id=admin.id,
        )
        base_b = KnowledgeBase(
            tenant_id=tenant_b.id,
            name="企业知识库",
            description="乙企业资料",
            created_by_user_id=other_admin.id,
        )
        db.add_all([base_a, base_b])
        db.commit()
        result = {
            "tenant_a": tenant_a.id,
            "tenant_b": tenant_b.id,
            "admin": admin.id,
            "manager": manager.id,
            "sales": sales.id,
            "other_admin": other_admin.id,
            "base_a": base_a.id,
            "base_b": base_b.id,
        }
    return result


def auth(seed: dict[str, object], role: str = "admin", other: bool = False) -> dict[str, str]:
    user_key = "other_admin" if other else role
    tenant_key = "tenant_b" if other else "tenant_a"
    token = create_access_token(
        user_id=seed[user_key],  # type: ignore[arg-type]
        tenant_id=seed[tenant_key],  # type: ignore[arg-type]
        role="admin" if other else role,
    )
    return {"Authorization": f"Bearer {token}"}


def upload_txt(
    client: TestClient,
    seed: dict[str, object],
    *,
    content: str = "用友 U8 Cloud 支持标准 API 对接、实施交付与价格咨询。" * 20,
    role: str = "admin",
) -> dict:
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed, role),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("product-guide.txt", content.encode(), "text/plain")},
    )
    assert response.status_code == 202, response.text
    return response.json()["data"]


def process_uploaded(factory: sessionmaker[Session], document_id: UUID) -> dict[str, object]:
    with factory() as db:
        document = db.get(KnowledgeDocument, document_id)
        job = db.scalar(
            select(KnowledgeProcessingJob)
            .where(KnowledgeProcessingJob.document_id == document_id)
            .order_by(KnowledgeProcessingJob.created_at.desc())
        )
        assert document is not None and job is not None
        tenant_id = document.tenant_id
        job_id = job.id
    return process_knowledge_document(
        document_id,
        tenant_id,
        job_id,
        provider=DeterministicTestEmbeddingProvider(1536),
        session_factory=factory,
    )


def upload_bytes(
    client: TestClient,
    seed: dict[str, object],
    filename: str,
    content: bytes,
    mime_type: str,
) -> dict:
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": (filename, content, mime_type)},
    )
    assert response.status_code == 202, response.text
    return response.json()["data"]


def supported_file(filename: str) -> tuple[bytes, str]:
    extension = filename.rsplit(".", 1)[-1]
    if extension in {"txt", "md", "csv"}:
        content = {
            "txt": "企业产品资料与标准 API 对接方案。" * 20,
            "md": "# API 对接\n\n" + "企业产品与实施交付说明。" * 20,
            "csv": "产品,说明\nU8 Cloud," + "标准接口对接方案" * 20,
        }[extension]
        mime = {"txt": "text/plain", "md": "text/markdown", "csv": "text/csv"}[extension]
        return content.encode(), mime
    output = BytesIO()
    if extension == "docx":
        document = Document()
        document.add_heading("交付方案", level=1)
        document.add_paragraph("企业产品与标准 API 对接说明。" * 20)
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    if extension == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "价格表"
        sheet.append(["产品", "说明"])
        sheet.append(["U8 Cloud", "标准 API 对接与交付方案" * 10])
        workbook.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "U8 Cloud API integration and delivery guide " * 12)
    output.write(pdf.tobytes())
    pdf.close()
    return output.getvalue(), "application/pdf"


def test_pgvector_schema_and_hnsw_index(postgres_client) -> None:
    _, factory = postgres_client
    with factory() as db:
        assert (
            db.scalar(text("SELECT extname FROM pg_extension WHERE extname='vector'")) == "vector"
        )
        index_definition = db.scalar(
            text(
                "SELECT indexdef FROM pg_indexes "
                "WHERE tablename='knowledge_chunks' AND indexdef ILIKE '%hnsw%'"
            )
        )
    assert "vector_cosine_ops" in index_definition


def test_knowledge_requires_authentication(postgres_client) -> None:
    client, _ = postgres_client
    assert client.get("/api/v1/knowledge/bases").status_code == 401


def test_tenant_registration_creates_default_knowledge_base(postgres_client) -> None:
    client, factory = postgres_client
    code = f"registered-{uuid4().hex[:8]}"
    response = client.post(
        "/api/v1/auth/register-tenant",
        json={
            "tenant_name": "新注册企业",
            "tenant_code": code,
            "admin_name": "管理员",
            "email": "registered@example.com",
            "password": "Test123456",
        },
    )
    assert response.status_code == 201
    with factory() as db:
        tenant = db.scalar(select(Tenant).where(Tenant.code == code))
        bases = list(db.scalars(select(KnowledgeBase).where(KnowledgeBase.tenant_id == tenant.id)))
    assert len(bases) == 1 and bases[0].name == "企业知识库"


def test_admin_can_create_update_and_delete_base(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    created = client.post(
        "/api/v1/knowledge/bases",
        headers=auth(seed),
        json={"name": "售前资料", "description": "产品与方案"},
    )
    assert created.status_code == 201
    base_id = created.json()["data"]["id"]
    updated = client.put(
        f"/api/v1/knowledge/bases/{base_id}",
        headers=auth(seed),
        json={"name": "售前标准资料"},
    )
    assert updated.json()["data"]["name"] == "售前标准资料"
    assert (
        client.delete(f"/api/v1/knowledge/bases/{base_id}", headers=auth(seed)).status_code == 200
    )


@pytest.mark.parametrize("role", ["manager", "sales"])
def test_non_admin_cannot_manage_bases(postgres_client, role: str) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.post(
        "/api/v1/knowledge/bases",
        headers=auth(seed, role),
        json={"name": "越权知识库"},
    )
    assert response.status_code == 403


def test_manager_cannot_delete_entire_base(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.delete(
        f"/api/v1/knowledge/bases/{seed['base_a']}", headers=auth(seed, "manager")
    )
    assert response.status_code == 403


def test_cross_tenant_base_is_hidden(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.get(f"/api/v1/knowledge/bases/{seed['base_b']}", headers=auth(seed))
    assert response.status_code == 404


@pytest.mark.parametrize("role", ["admin", "manager"])
def test_admin_and_manager_can_upload(postgres_client, role: str) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed, role=role)
    assert document["status"] == "uploaded"
    assert "storage_key" not in document and "tenant_id" not in document


def test_sales_cannot_upload(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed, "sales"),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("guide.txt", b"safe text", "text/plain")},
    )
    assert response.status_code == 403


@pytest.mark.parametrize(
    "filename", ["guide.md", "guide.docx", "prices.xlsx", "cases.csv", "guide.pdf"]
)
def test_supported_file_formats_upload_successfully(postgres_client, filename: str) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    content, mime_type = supported_file(filename)
    uploaded = upload_bytes(client, seed, filename, content, mime_type)
    assert uploaded["file_extension"] == filename.rsplit(".", 1)[-1]


def test_unsupported_file_type_returns_415(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("slides.pptx", b"unsafe", "application/octet-stream")},
    )
    assert response.status_code == 415
    assert response.json()["error_code"] == "UNSUPPORTED_FILE_TYPE"


def test_oversized_file_returns_413(postgres_client, monkeypatch: pytest.MonkeyPatch) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    monkeypatch.setattr(settings, "knowledge_max_file_size_mb", 0)
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("large.txt", b"one byte", "text/plain")},
    )
    assert response.status_code == 413


def test_tenant_id_form_injection_cannot_change_document_tenant(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    content = "企业资料和产品交付说明。" * 20
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={
            "knowledge_base_id": str(seed["base_a"]),
            "tenant_id": str(seed["tenant_b"]),
        },
        files={"file": ("safe.txt", content.encode(), "text/plain")},
    )
    assert response.status_code == 202
    with factory() as db:
        document = db.get(KnowledgeDocument, UUID(response.json()["data"]["id"]))
    assert document is not None and document.tenant_id == seed["tenant_a"]


def test_queue_failure_returns_clear_503_and_failed_status(
    postgres_client, monkeypatch: pytest.MonkeyPatch
) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)

    def unavailable(*_args):
        raise ConnectionError("redis unavailable")

    monkeypatch.setattr("app.api.v1.knowledge.process_document_task.delay", unavailable)
    content = "企业资料与标准交付方案。" * 20
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("queue-failure.txt", content.encode(), "text/plain")},
    )
    assert response.status_code == 503
    assert response.json()["error_code"] == "TASK_QUEUE_UNAVAILABLE"
    with factory() as db:
        document = db.scalar(
            select(KnowledgeDocument).where(
                KnowledgeDocument.original_filename == "queue-failure.txt"
            )
        )
    assert document is not None and document.status == DocumentStatus.FAILED


def test_fake_pdf_upload_is_rejected(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("fake.pdf", b"not really pdf", "application/pdf")},
    )
    assert response.status_code == 415
    assert response.json()["error_code"] == "INVALID_PDF_FORMAT"


def test_duplicate_file_in_same_base_is_rejected(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    upload_txt(client, seed)
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={
            "file": (
                "copy.txt",
                ("用友 U8 Cloud 支持标准 API 对接、实施交付与价格咨询。" * 20).encode(),
                "text/plain",
            )
        },
    )
    assert response.status_code == 409


def test_processing_creates_ready_chunks_and_is_idempotent(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    first = process_uploaded(factory, UUID(document["id"]))
    second = process_uploaded(factory, UUID(document["id"]))
    assert first["status"] == "ready"
    assert second["idempotent"] is True
    with factory() as db:
        saved = db.get(KnowledgeDocument, UUID(document["id"]))
        count = db.scalar(
            select(func.count())
            .select_from(KnowledgeChunk)
            .where(KnowledgeChunk.document_id == UUID(document["id"]))
        )
    assert saved is not None and saved.status == DocumentStatus.READY
    assert saved.chunk_count == count and count > 0


def test_search_returns_citation_and_writes_log(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    response = client.post(
        "/api/v1/knowledge/search",
        headers=auth(seed, "sales"),
        json={"query": "U8 Cloud API 对接", "top_k": 5, "min_score": 0},
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["results"] and data["results"][0]["document_name"] == "product-guide.txt"
    assert "citation_label" in data["results"][0]
    with factory() as db:
        assert db.scalar(select(func.count()).select_from(KnowledgeRetrievalLog)) == 1


def test_cross_tenant_search_and_document_access_are_isolated(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    other_headers = auth(seed, other=True)
    assert (
        client.get(
            f"/api/v1/knowledge/documents/{document['id']}", headers=other_headers
        ).status_code
        == 404
    )
    response = client.post(
        "/api/v1/knowledge/search",
        headers=other_headers,
        json={"query": "U8 Cloud", "top_k": 5, "min_score": 0},
    )
    assert response.status_code == 200
    assert response.json()["data"]["results"] == []


@pytest.mark.parametrize(
    ("method", "suffix"),
    [("get", "/download"), ("delete", ""), ("post", "/reprocess")],
)
def test_cross_tenant_document_mutations_are_hidden(
    postgres_client, method: str, suffix: str
) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    response = getattr(client, method)(
        f"/api/v1/knowledge/documents/{document['id']}{suffix}",
        headers=auth(seed, other=True),
    )
    assert response.status_code == 404


def test_cross_tenant_knowledge_base_search_returns_404(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    response = client.post(
        "/api/v1/knowledge/search",
        headers=auth(seed),
        json={"query": "产品资料", "knowledge_base_id": str(seed["base_b"])},
    )
    assert response.status_code == 404


def test_sales_only_sees_ready_documents(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    before = client.get("/api/v1/knowledge/documents", headers=auth(seed, "sales"))
    assert before.json()["data"]["total"] == 0
    process_uploaded(factory, UUID(document["id"]))
    after = client.get("/api/v1/knowledge/documents", headers=auth(seed, "sales"))
    assert after.json()["data"]["total"] == 1


def test_manager_can_view_chunks_but_sales_cannot(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    url = f"/api/v1/knowledge/documents/{document['id']}/chunks"
    assert client.get(url, headers=auth(seed, "manager")).status_code == 200
    assert client.get(url, headers=auth(seed, "sales")).status_code == 403


def test_disable_removes_document_from_retrieval_and_enable_restores_it(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    document_url = f"/api/v1/knowledge/documents/{document['id']}"
    assert client.patch(f"{document_url}/disable", headers=auth(seed)).status_code == 200
    hidden = client.post(
        "/api/v1/knowledge/search",
        headers=auth(seed),
        json={"query": "U8 Cloud", "min_score": 0},
    )
    assert hidden.json()["data"]["results"] == []
    assert client.patch(f"{document_url}/enable", headers=auth(seed)).status_code == 200
    visible = client.post(
        "/api/v1/knowledge/search",
        headers=auth(seed),
        json={"query": "U8 Cloud", "min_score": 0},
    )
    assert visible.json()["data"]["results"]


def test_min_score_filters_low_similarity_results(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    response = client.post(
        "/api/v1/knowledge/search",
        headers=auth(seed),
        json={"query": "完全无关的问题", "min_score": 1, "top_k": 5},
    )
    assert response.status_code == 200
    assert response.json()["data"]["results"] == []


def test_download_and_delete_remove_database_record_and_file(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    process_uploaded(factory, UUID(document["id"]))
    with factory() as db:
        stored = db.get(KnowledgeDocument, UUID(document["id"]))
        assert stored is not None
        stored_path = Path(settings.knowledge_storage_path) / stored.storage_key
        assert stored_path.is_file()
    download = client.get(
        f"/api/v1/knowledge/documents/{document['id']}/download", headers=auth(seed)
    )
    assert download.status_code == 200 and b"U8 Cloud" in download.content
    assert (
        client.delete(
            f"/api/v1/knowledge/documents/{document['id']}", headers=auth(seed)
        ).status_code
        == 200
    )
    with factory() as db:
        assert db.get(KnowledgeDocument, UUID(document["id"])) is None
        assert (
            db.scalar(
                select(func.count())
                .select_from(KnowledgeChunk)
                .where(KnowledgeChunk.document_id == UUID(document["id"]))
            )
            == 0
        )
    assert not stored_path.exists()


def test_reprocess_replaces_chunks_without_duplicates(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    document_id = UUID(document["id"])
    process_uploaded(factory, document_id)
    with factory() as db:
        first_count = db.scalar(
            select(func.count())
            .select_from(KnowledgeChunk)
            .where(KnowledgeChunk.document_id == document_id)
        )
    response = client.post(
        f"/api/v1/knowledge/documents/{document_id}/reprocess", headers=auth(seed)
    )
    assert response.status_code == 200
    process_uploaded(factory, document_id)
    with factory() as db:
        chunks = list(
            db.scalars(select(KnowledgeChunk).where(KnowledgeChunk.document_id == document_id))
        )
    assert len(chunks) == first_count
    assert len({chunk.chunk_index for chunk in chunks}) == first_count


def test_spreadsheet_chunk_keeps_sheet_and_row_metadata(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    content, mime_type = supported_file("prices.xlsx")
    document = upload_bytes(client, seed, "prices.xlsx", content, mime_type)
    process_uploaded(factory, UUID(document["id"]))
    with factory() as db:
        chunk = db.scalar(
            select(KnowledgeChunk).where(KnowledgeChunk.document_id == UUID(document["id"]))
        )
    assert chunk is not None
    assert chunk.sheet_name == "价格表" and chunk.row_start == 2


class WrongDimensionTaskProvider(EmbeddingProvider):
    name = "wrong_dimension"
    mode = "test"

    def embed_texts(self, texts: list[str]) -> list[list[float]]:
        return [[0.0] * 3 for _ in texts]


def test_dimension_mismatch_marks_processing_failed(postgres_client) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    document = upload_txt(client, seed)
    document_id = UUID(document["id"])
    with factory() as db:
        job = db.scalar(
            select(KnowledgeProcessingJob).where(KnowledgeProcessingJob.document_id == document_id)
        )
        assert job is not None
        job_id = job.id
    result = process_knowledge_document(
        document_id,
        seed["tenant_a"],  # type: ignore[arg-type]
        job_id,
        provider=WrongDimensionTaskProvider(1536),
        session_factory=factory,
    )
    assert result["error_code"] == "EMBEDDING_DIMENSION_MISMATCH"
    with factory() as db:
        failed = db.get(KnowledgeDocument, document_id)
    assert failed is not None and failed.status == DocumentStatus.FAILED


def test_scanned_pdf_processing_exposes_clear_failure(postgres_client, tmp_path: Path) -> None:
    client, factory = postgres_client
    seed = seed_enterprises(factory)
    path = tmp_path / "scan.pdf"
    pdf = fitz.open()
    pdf.new_page()
    pdf.save(path)
    pdf.close()
    response = client.post(
        "/api/v1/knowledge/documents/upload",
        headers=auth(seed),
        data={"knowledge_base_id": str(seed["base_a"])},
        files={"file": ("scan.pdf", path.read_bytes(), "application/pdf")},
    )
    assert response.status_code == 202
    result = process_uploaded(factory, UUID(response.json()["data"]["id"]))
    assert result["error_code"] == "OCR_REQUIRED"
    status_response = client.get(
        f"/api/v1/knowledge/documents/{response.json()['data']['id']}/status",
        headers=auth(seed),
    )
    assert status_response.json()["data"]["status"] == "failed"
